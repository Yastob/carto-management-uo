"""Réordonne data/collaborateurs.xlsx vers le nouvel ordre de colonnes attendu :
id, nom, prenom, poste, senior_manager_id (colonnes d'aide id_suggestion/doublon_id
déplacées après). Nécessaire une seule fois pour les fichiers créés avec l'ancien
gabarit (id, id_suggestion, doublon_id, nom, prenom, poste, senior_manager_id).

Ne touche ni rattachements.xlsx ni reunions.xlsx.
"""
import sys
import openpyxl
from template_builders import build_collaborateurs_workbook

PATH = sys.argv[1] if len(sys.argv) > 1 else "data/collaborateurs.xlsx"

wb = openpyxl.load_workbook(PATH, data_only=True)
ws = wb["Collaborateurs"]
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]

rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    r = dict(zip(headers, row))
    if not r.get("id"):
        continue
    rows.append(dict(
        id=r.get("id") or "",
        nom=r.get("nom") or "",
        prenom=r.get("prenom") or "",
        poste=r.get("poste") or "",
        senior_manager_id=r.get("senior_manager_id") or "",
    ))

build_collaborateurs_workbook(rows).save(PATH)
print(f"OK : {len(rows)} collaborateur(s) réordonné(s) dans {PATH}.")
