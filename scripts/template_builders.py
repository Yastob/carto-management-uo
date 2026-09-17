"""Fonctions partagées pour générer les 3 classeurs :
- collaborateurs.xlsx : référentiel (id, nom, prenom, poste, senior_manager_id), saisi à la
  main — le senior_manager_id y vit directement pour permettre de filtrer par périmètre de
  SM dès ce fichier chargé, sans dépendre de rattachements.xlsx.
- rattachements.xlsx  : Manager/CP/DP, compte_reference, tags — généré par editeur-rattachements.html.
- reunions.xlsx       : réunions de management + participants — généré par editeur-points.html.

Réutilisées par build_template.py (gabarits vierges avec exemples) et par
migrate_v1_to_v2.py (migration d'un ancien classeur organisation.xlsx à 1 fichier).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="170B8F")  # bleu marine mc2i
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF9E5")  # jaune pâle : à remplir
HELP_FILL = PatternFill("solid", fgColor="EDEDED")  # gris : aide, lecture seule
DOUBLON_FILL = PatternFill("solid", fgColor="FFC7CE")  # rouge pâle
EXAMPLE_FONT = Font(name=FONT_NAME, italic=True, color="808080")
HELP_FONT = Font(name=FONT_NAME, italic=True, color="808080", size=9)
NORMAL_FONT = Font(name=FONT_NAME)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14, color="170B8F")

N_COLLAB_ROWS = 150   # capacité Collaborateurs / Rattachements (lignes 2..151)
N_POINT_ROWS = 100    # capacité Points (lignes 2..101)

POSTES = ["Senior Manager", "Manager", "Chef de projet", "Directeur de projet", "Consultant"]

COLLAB_HEADERS = ["id", "id_suggestion", "doublon_id", "nom", "prenom", "poste", "senior_manager_id"]
COLLAB_WIDTHS = [8, 13, 11, 18, 18, 18, 16]

RATTACH_HEADERS = [
    "id", "manager_id", "chef_de_projet_id", "compte_reference",
    "tag_haut_potentiel", "tag_en_fragilite", "tag_consultant_isole", "actif", "date_maj",
]
RATTACH_WIDTHS = [8, 12, 18, 18, 18, 16, 20, 8, 12]

POINT_HEADERS = [
    "id", "nom", "animateur_id", "participants_ids", "type", "periodicite",
    "ordre_du_jour", "actif", "date_maj",
]
POINT_WIDTHS = [8, 30, 14, 30, 14, 16, 45, 8, 12]


def _style_header(ws, headers, widths):
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"


def _col(headers, name):
    return get_column_letter(headers.index(name) + 1)


def _lisezmoi(wb, lines):
    ws = wb.active
    ws.title = "Lisez-moi"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 100
    for i, (text, font) in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=text).font = font
    return ws


def build_collaborateurs_workbook(rows=None):
    """rows: liste de dicts avec les clés id/nom/prenom/poste/senior_manager_id (+ _example)."""
    rows = rows or []
    wb = openpyxl.Workbook()
    _lisezmoi(wb, [
        ("Collaborateurs — mode d'emploi", TITLE_FONT),
        ("", NORMAL_FONT),
        ("Référentiel : qui existe dans l'UO, et son vrai rattachement hiérarchique (Senior", NORMAL_FONT),
        ("Manager). C'est le SEUL fichier à saisir intégralement à la main — le reste", NORMAL_FONT),
        ("(Manager/CP opérationnels, tags, réunions) se fait depuis les pages HTML dédiées,", NORMAL_FONT),
        ("à partir de ce fichier.", NORMAL_FONT),
        ("", NORMAL_FONT),
        ("- id : identifiant unique (ex. C001). 'id_suggestion' te propose le prochain id", NORMAL_FONT),
        ("  disponible ; 'doublon_id' t'alerte si un id est utilisé deux fois.", NORMAL_FONT),
        ("- poste : Senior Manager / Manager / Chef de projet / Directeur de projet /", NORMAL_FONT),
        ("  Consultant. Un Chef de projet et un Directeur de projet sont traités de la même", NORMAL_FONT),
        ("  façon dans la visualisation — seul le poste affiché distingue les deux.", NORMAL_FONT),
        ("- senior_manager_id : le VRAI responsable hiérarchique (RH), quasi toujours renseigné", NORMAL_FONT),
        ("  sauf pour un Senior Manager lui-même. Choisis un id dans la liste, ou 'SM Sectoriel'", NORMAL_FONT),
        ("  si le vrai responsable est hors de cette UO (autre UO sectorielle). C'est ce champ", NORMAL_FONT),
        ("  qui permet de filtrer les pages d'édition sur le périmètre d'un seul SM.", NORMAL_FONT),
        ("", NORMAL_FONT),
        ("Mise à jour itérative : ajoute une ligne pour un nouvel arrivant, ne réécris pas", NORMAL_FONT),
        ("le fichier depuis zéro. Un départ se gère via 'actif=Non' dans rattachements.xlsx.", NORMAL_FONT),
        ("", NORMAL_FONT),
        ("Étape suivante : ouvre editeur-rattachements.html, charge ce fichier, et associe", NORMAL_FONT),
        ("à chaque collaborateur son Manager/CP opérationnel, ses tags, etc. → rattachements.xlsx.", NORMAL_FONT),
        ("Ouvre aussi editeur-points.html pour créer les réunions de management → reunions.xlsx.", NORMAL_FONT),
    ])

    ws = wb.create_sheet("Collaborateurs")
    _style_header(ws, COLLAB_HEADERS, COLLAB_WIDTHS)
    last_row = N_COLLAB_ROWS + 1

    for r, row in enumerate(rows, start=2):
        for name in ("id", "nom", "prenom", "poste", "senior_manager_id"):
            col = COLLAB_HEADERS.index(name) + 1
            cell = ws.cell(row=r, column=col, value=row.get(name, ""))
            cell.font = EXAMPLE_FONT if row.get("_example") else NORMAL_FONT

    id_col = _col(COLLAB_HEADERS, "id")
    poste_col = _col(COLLAB_HEADERS, "poste")
    sm_col = _col(COLLAB_HEADERS, "senior_manager_id")

    for r in range(2, last_row + 1):
        cell = ws.cell(row=r, column=COLLAB_HEADERS.index("id_suggestion") + 1)
        cell.value = '="C"&TEXT(ROW()-1,"000")'
        cell.font = HELP_FONT
        cell.fill = HELP_FILL
        cell = ws.cell(row=r, column=COLLAB_HEADERS.index("doublon_id") + 1)
        cell.value = f'=IF(AND(${id_col}{r}<>"",COUNTIF(${id_col}$2:${id_col}${last_row},${id_col}{r})>1),"⚠ Doublon","")'
        cell.font = HELP_FONT
        cell.fill = HELP_FILL
        for name in ("id", "nom", "prenom"):
            c2 = ws.cell(row=r, column=COLLAB_HEADERS.index(name) + 1)
            if not c2.value:
                c2.fill = INPUT_FILL
                c2.font = NORMAL_FONT

    ws.conditional_formatting.add(
        f"{_col(COLLAB_HEADERS,'doublon_id')}2:{_col(COLLAB_HEADERS,'doublon_id')}{last_row}",
        FormulaRule(formula=[f'${_col(COLLAB_HEADERS,"doublon_id")}2<>""'], fill=DOUBLON_FILL),
    )

    dv_poste = DataValidation(type="list", formula1=f'"{",".join(POSTES)}"', allow_blank=True)
    ws.add_data_validation(dv_poste)
    dv_poste.add(f"{poste_col}2:{poste_col}{last_row}")

    # Onglet caché "Listes" : combine le sentinel "SM Sectoriel" et tous les id de ce
    # même fichier, pour proposer une liste déroulante sur senior_manager_id sans avoir
    # à filtrer par poste (fragile en formule) — pas de contrainte de rôle sur ce choix.
    listes = wb.create_sheet("Listes")
    listes.sheet_view.showGridLines = False
    listes["A1"] = "SM Sectoriel"
    listes["A1"].font = NORMAL_FONT
    for i in range(2, last_row + 1):
        listes.cell(row=i, column=1, value=f"=Collaborateurs!{id_col}{i}").font = HELP_FONT
    listes.column_dimensions["A"].width = 20
    listes.sheet_state = "hidden"

    dv_sm = DataValidation(type="list", formula1=f"=Listes!$A$1:$A${last_row}", allow_blank=True)
    ws.add_data_validation(dv_sm)
    dv_sm.add(f"{sm_col}2:{sm_col}{last_row}")

    return wb


def build_rattachements_workbook(rows=None):
    """rows: liste de dicts avec les clés de RATTACH_HEADERS (+ _example optionnel).
    Normalement généré par editeur-rattachements.html — ce gabarit sert de secours
    pour une saisie/correction manuelle ponctuelle."""
    rows = rows or []
    wb = openpyxl.Workbook()
    _lisezmoi(wb, [
        ("Rattachements — mode d'emploi", TITLE_FONT),
        ("", NORMAL_FONT),
        ("Le VRAI responsable hiérarchique (senior_manager_id) est saisi directement dans", NORMAL_FONT),
        ("collaborateurs.xlsx. Ce fichier-ci ne gère que l'encadrement OPÉRATIONNEL :", NORMAL_FONT),
        ("Manager / Chef de projet / Directeur de projet, compte de référence et tags.", NORMAL_FONT),
        ("", NORMAL_FONT),
        ("Normalement généré et mis à jour depuis editeur-rattachements.html (à partir de", NORMAL_FONT),
        ("collaborateurs.xlsx), qui applique automatiquement les règles : aucun tag pour un", NORMAL_FONT),
        ("Senior Manager, 'Haut potentiel' réservé aux Consultants, 'Sectoriel' pour un", NORMAL_FONT),
        ("responsable hors UO. Une édition manuelle ici NE vérifie PAS ces règles", NORMAL_FONT),
        ("automatiquement — repasse par l'éditeur HTML en cas de doute.", NORMAL_FONT),
        ("", NORMAL_FONT),
        ("- id : doit correspondre exactement à un id de collaborateurs.xlsx.", NORMAL_FONT),
        ("- manager_id / chef_de_projet_id : id d'un autre collaborateur, ou 'Manager", NORMAL_FONT),
        ("  Sectoriel' / 'CP Sectoriel' si hors de cette UO. Peuvent pointer vers quelqu'un", NORMAL_FONT),
        ("  rattaché à un AUTRE Senior Manager — pas de contrainte de périmètre.", NORMAL_FONT),
        ("- compte_reference : client chez qui la personne travaille.", NORMAL_FONT),
        ("", NORMAL_FONT),
        ("Charge ensuite ce fichier avec collaborateurs.xlsx et reunions.xlsx dans index.html.", NORMAL_FONT),
    ])

    ws = wb.create_sheet("Rattachements")
    _style_header(ws, RATTACH_HEADERS, RATTACH_WIDTHS)
    last_row = N_COLLAB_ROWS + 1

    for r, row in enumerate(rows, start=2):
        for name in RATTACH_HEADERS:
            col = RATTACH_HEADERS.index(name) + 1
            cell = ws.cell(row=r, column=col, value=row.get(name, ""))
            cell.font = EXAMPLE_FONT if row.get("_example") else NORMAL_FONT

    for name in ["id", "manager_id", "chef_de_projet_id", "compte_reference", "date_maj"]:
        col = RATTACH_HEADERS.index(name) + 1
        for r in range(2, last_row + 1):
            c2 = ws.cell(row=r, column=col)
            if not c2.value:
                c2.fill = INPUT_FILL
                c2.font = NORMAL_FONT

    dv_oui_non = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    ws.add_data_validation(dv_oui_non)
    for name in ["tag_haut_potentiel", "tag_en_fragilite", "tag_consultant_isole", "actif"]:
        col = _col(RATTACH_HEADERS, name)
        dv_oui_non.add(f"{col}2:{col}{last_row}")

    return wb


def build_points_workbook(rows=None):
    """rows: liste de dicts avec les clés de POINT_HEADERS (+ _example optionnel).
    Normalement généré par editeur-points.html."""
    rows = rows or []
    wb = openpyxl.Workbook()
    _lisezmoi(wb, [
        ("Réunions de management — mode d'emploi", TITLE_FONT),
        ("", NORMAL_FONT),
        ("Ce fichier est normalement généré et mis à jour depuis editeur-points.html", NORMAL_FONT),
        ("(glisser-déposer de collaborateurs.xlsx + cases à cocher), pas besoin de le", NORMAL_FONT),
        ("remplir à la main. Tu peux aussi le modifier ici directement si besoin :", NORMAL_FONT),
        ("", NORMAL_FONT),
        ("- animateur_id et les id dans participants_ids doivent correspondre exactement", NORMAL_FONT),
        ("  aux id du fichier collaborateurs.xlsx (colonne 'id').", NORMAL_FONT),
        ("- participants_ids : plusieurs id séparés par une virgule (ex: C003,C004,C007).", NORMAL_FONT),
        ("  N'y remets pas l'animateur, il est déjà relié via animateur_id.", NORMAL_FONT),
        ("- Pour arrêter une réunion sans perdre l'historique, passe 'actif' à Non.", NORMAL_FONT),
        ("", NORMAL_FONT),
        ("Charge ensuite ce fichier avec collaborateurs.xlsx et rattachements.xlsx dans index.html.", NORMAL_FONT),
    ])

    ws = wb.create_sheet("Réunions")
    _style_header(ws, POINT_HEADERS, POINT_WIDTHS)
    last_row = N_POINT_ROWS + 1

    for r, row in enumerate(rows, start=2):
        for name in POINT_HEADERS:
            col = POINT_HEADERS.index(name) + 1
            cell = ws.cell(row=r, column=col, value=row.get(name, ""))
            cell.font = EXAMPLE_FONT if row.get("_example") else NORMAL_FONT

    for name in ["id", "nom", "animateur_id", "participants_ids", "ordre_du_jour", "date_maj"]:
        col = POINT_HEADERS.index(name) + 1
        for r in range(2, last_row + 1):
            c2 = ws.cell(row=r, column=col)
            if not c2.value:
                c2.fill = INPUT_FILL
                c2.font = NORMAL_FONT

    dv_type = DataValidation(type="list", formula1='"Equipe,Individuel,Autre"', allow_blank=True)
    ws.add_data_validation(dv_type)
    dv_type.add(f"{_col(POINT_HEADERS,'type')}2:{_col(POINT_HEADERS,'type')}{last_row}")

    dv_periodicite = DataValidation(
        type="list",
        formula1='"Hebdomadaire,2 fois par mois,Mensuel,Tous les 2 mois,Ponctuel,Autre"',
        allow_blank=True,
    )
    ws.add_data_validation(dv_periodicite)
    dv_periodicite.add(f"{_col(POINT_HEADERS,'periodicite')}2:{_col(POINT_HEADERS,'periodicite')}{last_row}")

    dv_actif = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    ws.add_data_validation(dv_actif)
    dv_actif.add(f"{_col(POINT_HEADERS,'actif')}2:{_col(POINT_HEADERS,'actif')}{last_row}")

    return wb
