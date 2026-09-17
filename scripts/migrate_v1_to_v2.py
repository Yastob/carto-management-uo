"""Migre l'ancien classeur à 1 fichier (organisation_test.xlsx, onglets
Collaborateurs/Points/Participations) vers la nouvelle structure à 3 fichiers :
data/collaborateurs.xlsx (référentiel minimal), data/rattachements.xlsx,
data/reunions.xlsx (participants fusionnés).

Applique au passage les nouvelles règles métier :
- un Senior Manager n'a aucun tag (les 3 tags sont remis à "Non")
- le tag 'Haut potentiel' est retiré des Manager / Chef de projet / Directeur de projet
- le tag 'Consultant détaché' est supprimé (n'existe plus)
"""
import sys
import openpyxl
from template_builders import (
    build_collaborateurs_workbook,
    build_rattachements_workbook,
    build_points_workbook,
)

SRC = sys.argv[1] if len(sys.argv) > 1 else "data/organisation_test.xlsx"

wb = openpyxl.load_workbook(SRC, data_only=True)


def rows_of(sheet_name):
    ws = wb[sheet_name]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        out.append(dict(zip(headers, row)))
    return out


old_collab = rows_of("Collaborateurs")
old_points = rows_of("Points")
old_particip = rows_of("Participations")

collab_rows = []
rattach_rows = []
notes = []

for r in old_collab:
    cid = r.get("id")
    poste = (r.get("poste") or "").strip()

    sm = (r.get("senior_manager_id") or "").strip() if r.get("senior_manager_id") else ""
    if sm == cid:
        notes.append(f"{cid} : senior_manager_id pointait vers lui-même — effacé, à corriger.")
        sm = ""

    collab_rows.append(dict(
        id=cid, nom=r.get("nom") or "", prenom=r.get("prenom") or "", poste=poste,
        senior_manager_id=sm,
    ))

    tag_hp = r.get("tag_haut_potentiel") or "Non"
    tag_fr = r.get("tag_en_fragilite") or "Non"
    tag_is = r.get("tag_consultant_isole") or "Non"
    # tag_consultant_detache existait avant : volontairement abandonné (supprimé du modèle).

    if poste == "Senior Manager" and (tag_hp == "Oui" or tag_fr == "Oui" or tag_is == "Oui"):
        notes.append(f"{cid} ({poste}) : tags remis à Non (un Senior Manager n'a aucun tag).")
        tag_hp, tag_fr, tag_is = "Non", "Non", "Non"
    elif poste in ("Manager", "Chef de projet", "Directeur de projet") and tag_hp == "Oui":
        notes.append(f"{cid} ({poste}) : tag 'Haut potentiel' retiré (réservé aux Consultants).")
        tag_hp = "Non"

    rattach_rows.append(dict(
        id=cid,
        manager_id=r.get("manager_id") or "",
        chef_de_projet_id=r.get("chef_de_projet_id") or "",
        compte_reference="",
        tag_haut_potentiel=tag_hp,
        tag_en_fragilite=tag_fr,
        tag_consultant_isole=tag_is,
        date_maj=r.get("date_maj") or "",
    ))

particip_by_point = {}
for p in old_particip:
    particip_by_point.setdefault(p["point_id"], []).append(p["collaborateur_id"])

point_rows = []
for p in old_points:
    point_rows.append(dict(
        id=p.get("id"),
        nom=p.get("nom") or "",
        animateur_id=p.get("animateur_id") or "",
        participants_ids=",".join(particip_by_point.get(p.get("id"), [])),
        type=p.get("type") or "",
        periodicite=p.get("periodicite") or "",
        ordre_du_jour=p.get("ordre_du_jour") or "",
        date_maj=p.get("date_maj") or "",
    ))

build_collaborateurs_workbook(collab_rows).save("data/collaborateurs.xlsx")
build_rattachements_workbook(rattach_rows).save("data/rattachements.xlsx")
build_points_workbook(point_rows).save("data/reunions.xlsx")

print(f"OK : {len(collab_rows)} collaborateurs, {len(point_rows)} points migrés.")
if notes:
    print("\nPoints d'attention :")
    for n in notes:
        print(" -", n)
